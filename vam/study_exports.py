"""Exports contain published source-grounded text; never execute model markup."""
from io import BytesIO
import textwrap
from vam.video_registry import DomainError


def clock(seconds, *, srt=False):
    value = max(0, round(float(seconds) * 1000))
    hours, value = divmod(value, 3600000)
    minutes, value = divmod(value, 60000)
    sec, ms = divmod(value, 1000)
    return f'{hours:02}:{minutes:02}:{sec:02}' + (f',{ms:03}' if srt else '')


def manuscript(records, format):
    if format == 'srt':
        return '\n\n'.join(f"{i}\n{clock(r['start_t'], srt=True)} --> {clock(r['end_t'], srt=True)}\n{r['text']}" for i, r in enumerate(records, 1)).encode('utf-8')
    return ('语音识别文稿（未经人工校对）\n\n' + '\n\n'.join(f"[{clock(r['start_t'])}] {r['text']}" for r in records)).encode('utf-8')


def export_study(artifact, format, max_pages=500):
    content = artifact['content']
    title, chapters = content['title'], content['chapters']
    warning = 'AI 生成内容，请结合原视频核对。' + (' 此产物来自旧版索引。' if artifact['stale'] else '')
    if format == 'md':
        lines = ['# ' + title, '', warning]
        for c in chapters:
            refs = [r for r in content.get('evidence', []) if r['evidence_id'] in c['source_ids']]
            lines += ['', f"## {clock(c['start_t'])} {c['title']}", *('- ' + p for p in c['points']),
                      '', '来源：' + '、'.join(f"{clock(r['start_t'])}-{clock(r['end_t'])}（{r['source']}）" for r in refs)]
        return '\n'.join(lines).encode('utf-8')
    # Use small text chunks to keep every exported page/slide readable, including
    # long Chinese words. Both formats preserve all points across continuation pages.
    pages = []
    for c in chapters:
        lines = []
        for point in c['points']:
            lines.extend(textwrap.wrap(point, 35, break_long_words=True, replace_whitespace=False) or [''])
        for offset in range(0, len(lines), 10):
            pages.append((c, lines[offset:offset + 10], offset > 0))
            if len(pages) > max_pages:
                raise DomainError(422, 'export_budget_exceeded', 'Courseware exceeds the configured page limit')
    output = BytesIO()
    if format == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for number, (c, lines, continued) in enumerate(pages, 1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for x, y, w, h, text, size, color in [
                (.65, .4, 12, .35, 'EviReel · 视频学习课件', 14, '64748B'),
                (.65, .9, 12, 1.2, '\n'.join(textwrap.wrap(c['title'] + ('（续）' if continued else ''), 40)), 22, '102344'),
                (.8, 2.2, 11.7, 4, '\n'.join(lines), 22, '253858'),
                (.65, 6.5, 12, .8, f"来源 {clock(c['start_t'])}–{clock(c['end_t'])} · 第 {number}/{len(pages)} 页\n{warning}", 11, '64748B')]:
                box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                tf = box.text_frame
                tf.word_wrap = True
                for index, line in enumerate(text.split('\n')):
                    p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.name, p.font.size, p.font.color.rgb = 'Microsoft YaHei', Pt(size), RGBColor.from_string(color)
        prs.core_properties.title = title
        prs.save(output)
    else:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        pdf = canvas.Canvas(output, pagesize=(960, 540))
        pdf.setTitle(title)
        for number, (c, lines, continued) in enumerate(pages, 1):
            pdf.setFillColorRGB(.03, .38, .94)
            pdf.rect(0, 525, 960, 15, fill=1, stroke=0)
            pdf.setFillColorRGB(.08, .15, .27)
            pdf.setFont('STSong-Light', 20)
            for index, title_line in enumerate(textwrap.wrap(c['title'] + ('（续）' if continued else ''), 40)):
                pdf.drawString(45, 475 - 25 * index, title_line)
            pdf.setFont('STSong-Light', 18)
            for index, line in enumerate(lines):
                pdf.drawString(55, 370 - 30 * index, line)
            pdf.setFont('STSong-Light', 10)
            pdf.drawString(45, 50, f"EviReel | 来源 {clock(c['start_t'])}–{clock(c['end_t'])} | {number}/{len(pages)}")
            pdf.drawString(45, 30, warning)
            pdf.showPage()
        pdf.save()
    return output.getvalue()
